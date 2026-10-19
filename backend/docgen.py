"""Generate questions from an uploaded document in the format the user asks for.

spec = {n, levels[]|None (balanced), difficulty|None, qtype: 'Short answer'|'Long answer'|'MCQ'|'Mixed', marks|None, topic|None, seed}
With an API key an LLM writes the questions from the document text. Without one, grounded templates are used
(key terms, headings and definitions found in the document), so the feature always works.
"""
import io, random, re, zipfile
from xml.sax.saxutils import escape

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer

import engine, llm

LEVELS = engine.LEVELS
QTYPES = ["Short answer", "Long answer", "MCQ", "Mixed"]
MAX_Q = 25
WEIGHTS = {"Remember": .15, "Understand": .25, "Apply": .25, "Analyze": .20, "Evaluate": .10, "Create": .05}
GENERIC = {"agenda", "introduction", "outline", "overview", "contents", "table of contents", "summary", "conclusion", "thank you", "thanks",
           "questions", "references", "q&a", "objectives", "learning objectives", "recap", "title", "bibliography", "acknowledgements",
           "comparison", "results", "demo", "demonstration", "applications", "examples", "example", "exercise", "exercises", "case study", "discussion",
           "next steps", "future work", "faq", "key points", "key takeaways", "takeaways", "review", "revision", "problem", "problems"}
NOISE_WORDS = {"slide", "figure", "example", "using", "used", "also", "can", "may", "one", "two", "page", "section", "shown", "following"}


# ----------------------------------------------------------------- understanding the document
def _clean_title(t):
    t = re.sub(r"^\s*(?:unit|chapter|module|section|lecture)?\s*\d+(?:\.\d+)*\s*[.):\-–]?\s*", "", t or "", flags=re.I).strip(" :-–.")
    n = len(t.split())
    return t if 1 <= n <= 7 and t.lower() not in GENERIC and re.search(r"[A-Za-z]{3}", t) else ""


DEF_RE = re.compile(r"^(?P<t>[A-Z][\w\-\*/ ]{1,50}?)\s+(?:is|are|refers to|is defined as|are defined as|means|can be defined as|is called|are called)\s+(?P<d>.{15,260})$")
COLON_RE = re.compile(r"^(?P<t>[A-Z][\w\-\* ]{1,40}):\s+(?P<d>.{15,260})$")
PRONOUN = {"this", "that", "it", "they", "these", "those", "there", "here", "which", "what", "we", "you", "he", "she", "such", "each", "some", "all"}


IMPERATIVE = {"mention", "note", "remember", "recall", "see", "consider", "discuss", "explain", "show", "use", "ask", "tell", "say", "refer", "review", "revise"}
BAD_DEF_START = {"only", "not", "also", "very", "always", "never", "both", "still", "just", "then", "however", "no", "yes"}


def _norm(t):
    t = re.sub(r"^(?:a|an|the)\s+", "", t.lower().strip())
    return re.sub(r"s$", "", re.sub(r"[\s\-]+", " ", t))


def _good_def(term, d, sentence=""):
    tw, dw = term.split(), d.split()
    plain_is = re.search(r"\b(?:is|are)\b", sentence) and not re.search(r"refers to|defined as|means|called", sentence)
    if plain_is and dw[0].lower() not in ("a", "an", "the"):        # "A* is optimal with ..." is a fact, not a definition
        return False
    return (tw[0].lower() not in PRONOUN | IMPERATIVE and len(dw) >= 5 and dw[0].lower() not in BAD_DEF_START
            and not any(w.lower() in ("is", "are", "was", "were", "be", "to") for w in tw))


def freq_phrases(texts, limit=40):
    """RAKE-style key phrases: frequent runs of content words (1-3 words) -> {phrase: (score, best_unit)}."""
    from collections import Counter, defaultdict
    from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS as SW
    cnt, where = Counter(), defaultdict(Counter)
    for i, t in enumerate(texts):
        for sent in re.split(r"[.!?\n;:()\[\]|•]+", t):
            toks = re.findall(r"[A-Za-z][A-Za-z\-\*]*", sent)
            for n in (1, 2, 3):
                for j in range(len(toks) - n + 1):
                    g = toks[j:j + n]
                    low = [w.lower() for w in g]
                    if low[0] in SW or low[-1] in SW or any(len(w) < 3 for w in g) or any(w in NOISE_WORDS or w in IMPERATIVE for w in low):
                        continue
                    if n > 1 and any(w in SW and w != "of" for w in low[1:-1]):
                        continue
                    k = " ".join(low)
                    cnt[k] += 1
                    where[k][i] += 1
    scored = {}
    for k, c in cnt.items():
        n = len(k.split())
        if (n == 1 and (c < 3 or len(k) < 6)) or (n > 1 and c < 2):
            continue
        scored[k] = c * (1 + 0.8 * (n - 1))
    keep = []
    for k in sorted(scored, key=scored.get, reverse=True):
        if not any(k in o or o in k for o in keep):              # drop sub/super-phrases of an already kept phrase
            keep.append(k)
        if len(keep) >= limit:
            break
    return {k: (scored[k], where[k].most_common(1)[0][0]) for k in keep}


def analyse(doc):
    """-> {terms (best first), defs, topics}; cached on the doc. Priority: defined terms > headings > frequent phrases."""
    if "analysis" in doc:
        return doc["analysis"]
    units = doc["units"]
    defs, seen = [], set()
    for i, u in enumerate(units):
        for s in re.split(r"(?<=[.!?])\s+|\n", u["text"]):
            s = s.strip(" •-–*\t")
            m = DEF_RE.match(s) or COLON_RE.match(s)
            if m and len(m["t"].split()) <= 6 and _good_def(m["t"].strip(), m["d"].strip(), s):
                term = re.sub(r"^(?:A|An|The)\s+", "", m["t"].strip())
                if _norm(term) in seen:
                    continue
                seen.add(_norm(term))
                defs.append({"term": term[:1].upper() + term[1:], "definition": m["d"].strip().rstrip("."), "ref": u["ref"], "unit": i})
    hi, lo, uniq = [], [], set()

    def add(lst, term, ref, unit):
        k = _norm(term)
        if k not in uniq:
            uniq.add(k)
            lst.append({"term": term, "ref": ref, "unit": unit})
    for d in defs:
        add(hi, d["term"], d["ref"], d["unit"])
    heads = []
    for i, u in enumerate(units):
        t = _clean_title(u["title"])
        if t:
            heads.append(t)
            add(hi, t, u["ref"], i)
    fp = freq_phrases([u["title"] + ". " + u["text"] for u in units])
    for k, (sc, i) in fp.items():
        add(lo if " " not in k else hi, k.title() if " " in k and len(k) < 30 else k, units[i]["ref"], i)
    topics = list(dict.fromkeys(heads)) or [t["term"] for t in hi if " " in t["term"]][:8] or [t["term"] for t in hi][:8]
    doc["analysis"] = {"hi": hi, "lo": lo, "terms": hi + lo, "defs": defs, "topics": topics[:8]}
    return doc["analysis"]


def overview(doc):
    a = analyse(doc)
    return {"topics": a["topics"][:6], "units": len(doc["units"]), "words": doc["words"], "definitions": len(a["defs"])}


def _sentence_with(term, unit_text):
    for s in re.split(r"(?<=[.!?])\s+|\n", unit_text):
        if term.lower() in s.lower() and 25 <= len(s) <= 300:
            return s.strip(" •-–*")
    return ""


# ----------------------------------------------------------------- planning
def distribute(n, levels):
    """How many questions at each Bloom level, ordered easy -> hard."""
    w = {l: (1.0 if levels else WEIGHTS[l]) for l in (levels or LEVELS)}
    tot = sum(w.values())
    raw = {l: n * v / tot for l, v in w.items()}
    cnt = {l: int(x) for l, x in raw.items()}
    for l in sorted(raw, key=lambda k: raw[k] - cnt[k], reverse=True)[: n - sum(cnt.values())]:
        cnt[l] += 1
    return [l for l in LEVELS if l in cnt for _ in range(cnt[l])]


def default_marks(qtype, level):
    if qtype == "MCQ":
        return 1
    if qtype == "Long answer":
        return 10
    return {"Remember": 2, "Understand": 3}.get(level, 5)


def qtype_for(spec_type, level):
    if spec_type != "Mixed":
        return spec_type
    return "Short answer" if LEVELS.index(level) <= 2 else "Long answer"


# ----------------------------------------------------------------- offline generation
OFFLINE = {
    "Remember": ["Define {t}.", "State the purpose of {t}.", "List the key characteristics of {t}.", "Identify the main components of {t}."],
    "Understand": ["Explain the concept of {t} in your own words.", "Describe how {t} works.", "Summarize the main idea behind {t}.", "Discuss the significance of {t}."],
    "Apply": ["Apply {t} to a real-world scenario of your choice, showing the steps involved.", "Demonstrate how {t} can be used to solve a simple problem.", "Illustrate {t} with a worked example."],
    "Analyze": ["Analyze the strengths and weaknesses of {t}.", "Compare {t} with {u}.", "Differentiate between {t} and {u}.", "Examine how {t} relates to {u}."],
    "Evaluate": ["Evaluate the effectiveness of {t} for practical problems, and justify your conclusion.", "Critique the use of {t}, considering {u}.", "Assess the advantages and limitations of {t}."],
    "Create": ["Design a solution that makes use of {t}.", "Propose an improvement to {t} and justify it.", "Develop a plan to apply {t} to a new problem."],
}


def _trim(s, n=28):
    w = s.split()
    return " ".join(w[:n]) + ("…" if len(w) > n else "")


def _offline_mcq(doc, count, rnd):
    defs = analyse(doc)["defs"]
    if len(defs) < 4:
        return [], "Not enough clear definitions in the document to build multiple-choice questions without an API key."
    pool = defs[:]
    rnd.shuffle(pool)
    out = []
    for d in pool[:count]:
        others = [x for x in defs if x["term"] != d["term"]]
        rnd.shuffle(others)
        opts = [_trim(d["definition"])] + [_trim(x["definition"]) for x in others[:3]]
        order = list(range(4))
        rnd.shuffle(order)
        shuffled = [opts[i] for i in order]
        out.append({"question": f"Which of the following best describes {d['term']}?", "type": "MCQ", "level": "Remember",
                    "options": shuffled, "answer": "ABCD"[order.index(0)] + ") " + opts[0], "source": d["ref"],
                    "reference": d["definition"][:220]})
    return out, ("" if len(out) >= count else f"Only {len(out)} multiple-choice question(s) could be built from the definitions found.")


def offline_generate(doc, spec):
    rnd = random.Random(f"{doc['id']}-{spec.get('seed', 0)}-{spec['n']}")
    a = analyse(doc)
    hi, lo = a["hi"][:], a["lo"][:]
    rnd.shuffle(hi)
    rnd.shuffle(lo)
    terms = hi + lo
    if spec.get("topic"):
        tw = {w for w in re.findall(r"[a-z]{3,}", spec["topic"].lower())}
        hit = [t for t in terms if tw & set(re.findall(r"[a-z]{3,}", t["term"].lower())) or tw & set(re.findall(r"[a-z]{3,}", doc["units"][t["unit"]]["text"].lower()))]
        terms = hit or terms
        hi = [t for t in terms if t in hi] or terms
    if not terms:
        return [], ["I couldn't find enough key terms in this document to write questions from."]
    notes, questions = [], []
    plan = distribute(spec["n"], spec["levels"])
    if spec["qtype"] == "MCQ":
        qs, note = _offline_mcq(doc, spec["n"], rnd)
        questions += qs
        notes += [note] if note else []
        if spec["levels"] and any(l not in ("Remember", "Understand") for l in spec["levels"]):
            notes.append("Without an API key, multiple-choice questions are recall-level only. Add a GEMINI_API_KEY for higher-level MCQs.")
        plan = []
    used, ti = set(), 0
    for lvl in plan:
        for attempt in range(8):
            t = terms[ti % len(terms)]
            ti += 1
            pool = [x for x in hi if x["term"] != t["term"]] or terms
            other = next((x for x in pool if x["unit"] == t["unit"]), None) or pool[(ti + 2) % len(pool)]
            tmpl = rnd.choice(OFFLINE[lvl])
            q = tmpl.format(t=t["term"], u=other["term"])
            if (q.lower() in used) or (tmpl.count("{u}") and other["term"].lower() == t["term"].lower()):
                continue
            if engine.predict(q)["level"] != lvl and attempt < 5:       # keep the wording the model agrees with
                continue
            used.add(q.lower())
            ref = _sentence_with(t["term"], doc["units"][t["unit"]]["text"])
            questions.append({"question": q, "type": qtype_for(spec["qtype"], lvl), "level": lvl, "options": None, "answer": None,
                              "source": t["ref"], "reference": ref})
            break
    notes.append("Written without an API key from the key terms and headings in your document. Add a GEMINI_API_KEY for questions that use the document's actual facts.")
    return questions, notes


# ----------------------------------------------------------------- LLM generation
def _excerpt(doc, spec, budget=14000):
    units = doc["units"]
    if spec.get("topic"):
        tw = set(re.findall(r"[a-z]{3,}", spec["topic"].lower()))
        scored = sorted(range(len(units)), key=lambda i: -len(tw & set(re.findall(r"[a-z]{3,}", (units[i]["title"] + " " + units[i]["text"]).lower()))))
        units = [units[i] for i in sorted(scored[:12])]
    per = max(350, min(1400, budget // max(1, len(units))))
    out, used = [], 0
    for u in units:
        chunk = f"[{u['ref']}] {u['title']}\n{u['text'][:per]}"
        if used + len(chunk) > budget:
            break
        out.append(chunk)
        used += len(chunk)
    return "\n\n".join(out)


def llm_generate(doc, spec):
    if not llm.available():
        return None, None
    verbs = {l: engine._df[engine._df.level == l].question.str.split().str[0].str.lower().value_counts().head(7).index.tolist() for l in LEVELS}
    plan = distribute(spec["n"], spec["levels"])
    counts = ", ".join(f"{plan.count(l)} x {l}" for l in LEVELS if l in plan)
    qtype = {"Short answer": "short-answer questions", "Long answer": "long-answer / essay questions", "MCQ": "multiple-choice questions (MCQ)",
             "Mixed": "a mix of MCQ, short-answer and long-answer questions"}[spec["qtype"]]
    material = _excerpt(doc, spec)
    questions, seen, notes = [], set(), []
    for attempt in range(2):
        need = spec["n"] - len(questions) + (2 if attempt == 0 else 0)
        ai = llm.ask_json(f"""You are an expert exam setter. Using ONLY the study material below, write {need} {qtype}.
Required Bloom levels: {counts or 'balanced'}. Difficulty: {spec.get('difficulty') or 'mixed'}. Marks per question: {spec.get('marks') or 'choose sensibly (MCQ 1, short 2-5, long 8-10)'}.
{('Focus on this topic: ' + spec['topic']) if spec.get('topic') else ''}

STYLE RULES
- Written questions are ONE imperative sentence that STARTS with an action verb that fits the level. Verbs by level: {'; '.join(f"{l}: {', '.join(v)}" for l, v in verbs.items())}.
- Never start a written question with What, Which, How, Why, In, During, Given or Consider.
- MCQs: a clear stem plus exactly 4 options and the correct letter. The stem may start with "Which" or "What".
- Every question must be answerable from the material. Set "source" to the [slide/page/section] tag it came from.
- "answer" is a concise model answer (max 40 words) taken from the material (for MCQ: the letter and the option text).
{('Do NOT repeat: ' + '; '.join(q['question'] for q in questions)) if questions else ''}

STUDY MATERIAL
{material}

Return JSON: {{"questions": [{{"question": "...", "type": "MCQ|Short answer|Long answer", "level": "...", "difficulty": "Easy|Medium|Hard", "marks": 0, "source": "...", "options": ["..."], "answer": "..."}}]}}""")
        if not ai:
            break
        for it in ai.get("questions", []):
            t = it.get("type") if it.get("type") in QTYPES else ("MCQ" if it.get("options") else "Short answer")
            q = str(it.get("question", "")).strip()
            if not q or q.lower() in seen:
                continue
            if t == "MCQ":
                opts = [str(o).strip() for o in (it.get("options") or [])]
                if len(opts) != 4 or len(q.split()) < 5:
                    continue
                lvl = it.get("level") if it.get("level") in LEVELS else "Remember"
                rec = {"question": q, "type": "MCQ", "level": lvl, "options": opts, "answer": str(it.get("answer", "")).strip()}
            else:
                q = engine._valid_generated({"question": q}, seen, min_words=3)
                if not q:
                    continue
                lvl = engine.predict(q)["level"]
                rec = {"question": q, "type": t, "level": lvl, "options": None, "answer": str(it.get("answer", "")).strip() or None}
            seen.add(q.lower())
            rec.update({"source": str(it.get("source", "")).strip("[] "), "reference": "",
                        "difficulty_ai": it.get("difficulty") if it.get("difficulty") in engine.DIFFS else None, "marks_ai": it.get("marks")})
            questions.append(rec)
        if len(questions) >= spec["n"]:
            break
    if spec["levels"]:                                   # prefer the requested levels, keep order easy -> hard
        questions.sort(key=lambda r: r["level"] not in spec["levels"])
    return questions[: spec["n"]], notes


# ----------------------------------------------------------------- main entry
def generate(doc, spec):
    spec = {"n": 10, "levels": None, "difficulty": None, "qtype": "Short answer", "marks": None, "topic": None, "seed": 0, **spec}
    spec["n"] = max(1, min(MAX_Q, int(spec["n"])))
    questions, notes, source = None, [], "llm"
    questions, extra = llm_generate(doc, spec)
    if questions is None or len(questions) == 0:
        source = "offline"
        if llm.available():
            notes.append("The AI could not be reached, so these were written from the key terms in your document. " + (llm.last_error or ""))
        questions, extra = offline_generate(doc, spec)
    else:
        notes += extra or []
        if len(questions) < spec["n"]:
            more, more_notes = offline_generate(doc, {**spec, "n": spec["n"] - len(questions)})
            questions += more
            source, notes = "llm+offline", notes + [f"The AI returned {len(questions) - len(more)} valid question(s); the rest were filled from the document's key terms."]
    if source == "offline":
        notes += extra
    questions.sort(key=lambda r: LEVELS.index(r["level"]))
    out = []
    for i, r in enumerate(questions[: spec["n"]], 1):
        d = spec["difficulty"] or r.pop("difficulty_ai", None) or engine.predict(r["question"])["difficulty"]
        r.pop("difficulty_ai", None)
        marks = spec["marks"] or (r.pop("marks_ai", None) if isinstance(r.get("marks_ai"), int) and 0 < r.get("marks_ai") <= 20 else None) or default_marks(r["type"], r["level"])
        r.pop("marks_ai", None)
        out.append({**r, "no": i, "difficulty": d, "marks": marks})
    return {"title": f"Questions from {doc['name']}", "filename": doc["name"], "doc_id": doc["id"], "questions": out,
            "total_marks": sum(q["marks"] for q in out), "spec": spec, "source": source, "notes": notes}


def spec_text(spec):
    lv = ", ".join(spec["levels"]) if spec.get("levels") else "balanced levels"
    bits = [f"{spec['n']} questions", spec["qtype"], lv, (spec.get("difficulty") or "any difficulty")]
    if spec.get("marks"):
        bits.append(f"{spec['marks']} marks each")
    if spec.get("topic"):
        bits.append(f"topic: {spec['topic']}")
    return " · ".join(bits)


# ----------------------------------------------------------------- export: Word + Excel
def _p(text, bold=False, size=None, indent=0, page_break=False, keep=False):
    rpr = ("<w:b/>" if bold else "") + (f'<w:sz w:val="{size}"/>' if size else "")
    ppr = (f'<w:ind w:left="{indent}"/>' if indent else "") + ('<w:keepNext/>' if keep else "")
    br = '<w:r><w:br w:type="page"/></w:r>' if page_break else ""
    return f'<w:p><w:pPr>{ppr}</w:pPr>{br}<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def to_docx(d, include_answers=True):
    """A plain, valid .docx question paper (no extra libraries)."""
    body = [_p(d["title"], True, 34), _p(f"Total marks: {d['total_marks']}      Time: ________", False, 22), _p("")]
    for q in d["questions"]:
        body.append(_p(f"{q['no']}. {q['question']}   [{q['marks']} mark{'s' if q['marks'] != 1 else ''}]  ({q['level']})", True, keep=True))
        for i, o in enumerate(q.get("options") or []):
            body.append(_p(f"{'ABCD'[i]}) {o}", indent=480))
        body.append(_p(""))
    ans = [q for q in d["questions"] if q.get("answer") or q.get("reference")]
    if include_answers and ans:
        body.append(_p("Answer key", True, 30, page_break=True))
        for q in ans:
            body.append(_p(f"{q['no']}. {q.get('answer') or q['reference']}" + (f"   (Source: {q['source']})" if q.get("source") else "")))
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           "<w:body>" + "".join(body) + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr></w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", doc)
    buf.seek(0)
    return buf


def to_xlsx(d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    head = ["No.", "Question", "Type", "Bloom level", "Difficulty", "Marks", "Source", "Options", "Answer / reference"]
    for c, h in enumerate(head, 1):
        x = ws.cell(row=1, column=c, value=h)
        x.fill, x.font = PatternFill("solid", fgColor="1D4ED8"), Font(bold=True, color="FFFFFF")
    for r, q in enumerate(d["questions"], 2):
        vals = [q["no"], q["question"], q["type"], q["level"], q["difficulty"], q["marks"], q.get("source", ""),
                "\n".join(f"{'ABCD'[i]}) {o}" for i, o in enumerate(q.get("options") or [])), q.get("answer") or q.get("reference") or ""]
        for c, v in enumerate(vals, 1):
            x = ws.cell(row=r, column=c, value=v)
            if isinstance(v, str):
                x.data_type = "s"
            x.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHI", [6, 70, 14, 14, 12, 8, 14, 50, 50]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.cell(row=len(d["questions"]) + 3, column=1, value=f"Total marks: {d['total_marks']}")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
