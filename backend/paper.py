"""Question-paper PDF  ->  individual questions  ->  Bloom level for each.

Pipeline: extract text -> drop headers/footers -> split into numbered questions and sub-parts
-> strip marks and [CO/BL] tags -> predict Bloom level + difficulty -> summarise the paper.
Everything runs locally (no API key needed).
"""
import io, os, re, subprocess, tempfile
from collections import Counter

import engine

LEVELS = engine.LEVELS
MAX_PAGES, MAX_Q = 80, 150
LEVEL_BY_NUM = {str(i + 1): l for i, l in enumerate(LEVELS)}      # paper tags: L1..L6 / BL1..BL6 / K1..K6


# ----------------------------------------------------------------- 1. text extraction
def extract_pages(data):
    """Return (pages_text, total_page_count). Tries PyMuPDF, then pypdf, then the pdftotext tool."""
    try:
        try:
            import pymupdf as fz
        except ImportError:
            import fitz as fz
        doc = fz.open(stream=data, filetype="pdf")
        return [doc[i].get_text() for i in range(min(doc.page_count, MAX_PAGES))], doc.page_count
    except ImportError:
        pass
    except Exception as e:
        raise ValueError(f"This file could not be opened as a PDF ({str(e)[:80]}).")
    try:
        from pypdf import PdfReader
        rd = PdfReader(io.BytesIO(data))
        if rd.is_encrypted:
            raise ValueError("This PDF is password-protected. Remove the password and upload it again.")
        return [(rd.pages[i].extract_text() or "") for i in range(min(len(rd.pages), MAX_PAGES))], len(rd.pages)
    except ImportError:
        pass
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"This file could not be opened as a PDF ({str(e)[:80]}).")
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as t:
            t.write(data)
        out = subprocess.run(["pdftotext", "-layout", t.name, "-"], capture_output=True, text=True, timeout=60).stdout
        os.unlink(t.name)
        pages = out.split("\f")
        return pages[:MAX_PAGES], len(pages)
    except Exception:
        raise ValueError("No PDF reader is installed on the server. Run: pip install pymupdf")


_norm = lambda s: re.sub(r"\d+", "#", s.lower())
PAGE_NO = re.compile(r"^\s*(?:page\s*)?\d{1,3}(?:\s*(?:of|/)\s*\d{1,3})?\s*$", re.I)


def clean_lines(pages):
    per = [[l.strip() for l in p.splitlines() if l.strip()] for p in pages]
    rep = set()
    if len(per) >= 3:                                   # lines repeated on most pages = running header/footer
        cnt = Counter(_norm(l) for pg in per for l in set(pg) if len(l) >= 8 and sum(c.isalpha() for c in l) >= 5)
        rep = {k for k, v in cnt.items() if v >= max(3, int(0.6 * len(per)))}
    out = []
    for pg in per:
        pg = [l for l in pg if _norm(l) not in rep]
        while pg and PAGE_NO.match(pg[-1]):             # trailing page number
            pg.pop()
        while pg and PAGE_NO.match(pg[0]):
            pg.pop(0)
        out.extend(pg)
    return out


# ----------------------------------------------------------------- 2. splitting into questions
Q_RE = re.compile(r"^(?:(?:Q(?:uestion)?|Qn)\.?\s*(\d{1,2})\s*[.):\-]?|(\d{1,2})\s*[.):])(?!\d)\s*(.*)$", re.I)
Q_RE_OCR = re.compile(r"^(?:(?:Q(?:uestion)?|Qn)\.?\s*(\d{1,2})\s*[.):,;\-]?|(\d{1,2})\s*[.):,;])(?!\d)\s*(.*)$", re.I)   # scans: "5," for "5."
SUB_RE = re.compile(r"^\(?([a-h]|[ivx]{1,4})\)\s*(.*)$", re.I)
SECTION_RE = re.compile(r"^(?:section|part|unit|module)\s*[-–:]?\s*[A-Za-z0-9IVX]+\b", re.I)
OR_RE = re.compile(r"^(?:or|\(or\)|[-–—]\s*or\s*[-–—])$", re.I)
NOISE_RE = re.compile(
    r"^(?:(?:answer|attempt)\s+(?:all|any)\b|(?:instructions?|notes?)\s*[:\-]|time\s*[:\-]|duration\b|max(?:imum)?\.?\s*marks|total\s+marks|"
    r"roll\s*no|reg(?:ister)?\.?\s*no|hall\s*ticket|all\s+questions\s+carry|each\s+question\s+carries|marks\s*:|\(?\s*\d+\s*[x×]\s*\d+\s*=\s*\d+)", re.I)

TAG_BODY = re.compile(r"(?:\s*(?:C[O0]\s*-?\s*[\dlIt]+|PO\s*-?\s*\d+|(?:BL|BTL|RBT|L|K)\s*-?\s*[1-6]|\d{1,2}\s*(?:marks?|m|pts?)?)\s*[,;/&+]?)+", re.I)
LEVEL_TAG = re.compile(r"\b(?:BL|BTL|RBT|L|K)\s*-?\s*([1-6])\b", re.I)


def strip_meta(text):
    """Remove [CO2, L3], (10 marks) and similar. Returns (clean_text, marks, declared_bloom_level)."""
    marks, declared = None, None

    def bracket(m):
        nonlocal marks, declared
        inner = m.group(1)
        if not TAG_BODY.fullmatch(inner):
            return m.group(0)
        at_end = text[m.end():].strip(" .;:") == ""
        only_int = re.fullmatch(r"\s*\d{1,2}\s*", inner)
        if only_int and not at_end:
            return m.group(0)
        lv = LEVEL_TAG.search(inner)
        if lv:
            declared = LEVEL_BY_NUM[lv.group(1)]
        mk = re.search(r"(\d{1,2})\s*(?:marks?|m|pts?)\b", inner, re.I) or (re.fullmatch(r"\s*(\d{1,2})\s*", inner) if only_int else None)
        if mk and marks is None:
            marks = int(mk.group(1))
        return " "

    t = re.sub(r"[\[({|]([^\[\](){}|]{1,30})[\])}|]", bracket, text)       # OCR often swaps bracket types
    m = re.search(r"[\s\-–—,]*\b(\d{1,2})\s*(?:marks?|m)\b\s*[.)\]]*\s*$", t, re.I)       # "... 5 Marks" at the end
    if m:
        marks = marks or int(m.group(1))
        t = t[:m.start()]
    t = re.sub(r"\s+", " ", t).strip()
    return re.sub(r"\s+([.,;:?])", r"\1", t), marks, declared


def _join(lines):
    out = ""
    for ln in lines:
        if out.endswith("-") and ln[:1].islower():
            out = out[:-1] + ln
        else:
            out = (out + " " + ln).strip()
    return out


def _accept(num, last):
    return num == last + 1 or num == 1 or (last == 0 and num <= 3)


def split_numbered(lines, ocr=False):
    qre = Q_RE_OCR if ocr else Q_RE
    items, cur, sub, last, closed, section = [], None, None, 0, True, ""
    for s in lines:
        if OR_RE.match(s):
            continue
        if SECTION_RE.match(s) and len(s) < 70:
            section, closed = s, True
            continue
        if NOISE_RE.match(s):
            closed = True
            continue
        m = qre.match(s)
        if m:
            num = int(m.group(1) or m.group(2))
            if _accept(num, last):
                cur = {"num": str(num), "stem": [], "subs": [], "section": section}
                items.append(cur)
                sub, last, closed = None, num, False
                rest = m.group(3) or ""
                sm = SUB_RE.match(rest) if rest else None
                if sm:                                  # "6. (a) Apply ..." - sub-part label on the same line as the number
                    sub = {"label": sm.group(1).lower(), "lines": [sm.group(2)] if sm.group(2) else []}
                    cur["subs"].append(sub)
                elif rest:
                    cur["stem"].append(rest)
                continue
        if cur is None or closed:
            continue
        sm = SUB_RE.match(s)
        if sm:
            sub = {"label": sm.group(1).lower(), "lines": [sm.group(2)] if sm.group(2) else []}
            cur["subs"].append(sub)
        elif sub is not None:
            sub["lines"].append(s)
        else:
            cur["stem"].append(s)
    return items


def split_by_verbs(lines):
    """Fallback for papers without question numbers: a new question starts at a line beginning with an action verb."""
    items, cur = [], None
    for s in lines:
        fw = re.sub(r"[^a-z]", "", s.split()[0].lower()) if s.split() else ""
        if fw in engine.ALLOWED_START and fw not in engine.BANNED_START and len(s.split()) >= 4 and not NOISE_RE.match(s):
            cur = {"num": str(len(items) + 1), "stem": [s], "subs": [], "section": ""}
            items.append(cur)
        elif cur and not re.search(r"[.?]$", _join(cur["stem"])) and not NOISE_RE.match(s) and not SECTION_RE.match(s):
            cur["stem"].append(s)
    return items


def is_mcq(stem, subs):
    if not (3 <= len(subs) <= 6) or len(stem.split()) < 4:
        return False
    avg = sum(len(t.split()) for _, t in subs) / len(subs)
    verbs = sum(1 for _, t in subs if t.split() and t.split()[0].lower() in engine.ALLOWED_START)
    return avg <= 10 and verbs < len(subs) / 2


# ----------------------------------------------------------------- 3. predict + summarise
def build_records(items):
    recs = []
    for it in items:
        stem = _join(it["stem"])
        subs = [(s["label"], _join(s["lines"])) for s in it["subs"] if s["lines"]]
        rows = []
        if subs and is_mcq(stem, subs):
            rows.append((it["num"], "", stem, "", [t for _, t in subs]))
        elif subs:
            rows += [(it["num"], lab, t, stem, None) for lab, t in subs]
        elif stem:
            rows.append((it["num"], "", stem, "", None))
        for num, part, text, ctx, opts in rows:
            clean, marks, declared = strip_meta(text)
            ctx_clean = strip_meta(ctx)[0] if ctx else ""
            if len(clean.split()) < 3:
                continue
            ptext = clean if len(clean.split()) >= 6 or not ctx_clean else ctx_clean + " " + clean
            p = engine.predict(ptext)
            ranked = sorted(p["level_probs"].items(), key=lambda kv: -kv[1])
            alt = ranked[1] if len(ranked) > 1 and ranked[1][1] >= 0.25 else None
            recs.append({
                "id": f"{num}({part})" if part else num, "number": num, "part": part, "section": it["section"],
                "text": clean, "context": ctx_clean[:220], "marks": marks, "type": "MCQ" if opts else "Written",
                "options": len(opts) if opts else 0,
                "level": p["level"], "level_confidence": p["level_confidence"], "difficulty": p["difficulty"],
                "alt_level": alt[0] if alt else None, "alt_confidence": round(alt[1], 3) if alt else None,
                "next_level": p["next_level"], "declared_level": declared,
                "low_confidence": p["level_confidence"] < 0.5,
            })
    return recs[:MAX_Q]


def summarize(recs):
    n = len(recs)
    counts, marks = {l: 0 for l in LEVELS}, {l: 0 for l in LEVELS}
    for r in recs:
        counts[r["level"]] += 1
        marks[r["level"]] += r["marks"] or 0
    total_marks = sum(marks.values())
    by_level = [{"level": l, "count": counts[l], "percent": round(100 * counts[l] / n) if n else 0, "marks": marks[l]} for l in LEVELS]
    diff = Counter(r["difficulty"] for r in recs)
    low = counts["Remember"] + counts["Understand"]
    high = n - low
    tagged = [r for r in recs if r["declared_level"]]
    mismatch = [r for r in tagged if r["declared_level"] != r["level"]]
    insights = []
    if n:
        top = max(LEVELS, key=lambda l: counts[l])
        if counts[top] / n >= 0.5:
            insights.append(f"{round(100 * counts[top] / n)}% of the questions are at the {top} level.")
        insights.append(f"Lower-order (Remember + Understand): {round(100 * low / n)}%. Higher-order (Apply and above): {round(100 * high / n)}%.")
        if low / n > 0.6:
            insights.append("The paper leans toward recall and explanation; consider adding Apply, Analyze or Evaluate questions.")
        missing = [l for l in LEVELS if counts[l] == 0]
        if missing and n >= 8:
            insights.append("No questions at: " + ", ".join(missing) + ".")
        if total_marks:
            hm = sum(marks[l] for l in LEVELS[2:])
            insights.append(f"{round(100 * hm / total_marks)}% of the detected marks ({hm} of {total_marks}) go to Apply-level questions or higher.")
        if tagged:
            insights.append(f"{len(tagged) - len(mismatch)} of {len(tagged)} questions with a Bloom tag in the paper agree with the prediction.")
        lc = sum(r["low_confidence"] for r in recs)
        if lc:
            insights.append(f"{lc} prediction(s) have low confidence; please review them manually.")
    return {"total": n, "by_level": by_level, "difficulty": {d: diff.get(d, 0) for d in engine.DIFFS},
            "total_marks": total_marks or None, "lower_order_pct": round(100 * low / n) if n else 0,
            "higher_order_pct": round(100 * high / n) if n else 0, "tagged": len(tagged), "mismatches": len(mismatch),
            "insights": insights}


def _ocr_fix(lines):
    """Typical OCR slips on question numbers: 'l.' / 'I.' / '|.' for '1.', and '1 )' spacing."""
    out, fixed = [], False
    for s in lines:
        if not fixed and re.match(r"^[lI|!]\s*[.)]\s+\S", s):
            s, fixed = "1." + s[s.index(re.search(r"[.)]", s).group(0)) + 1:], True
        out.append(re.sub(r"^(\d{1,2})\s+([.)])", r"\1\2", s))
    return out


def analyze_doc(doc):
    """Question-paper mode for an ingested document (digital PDF, scan, photo, Word or PowerPoint)."""
    lines, notes = doc["lines"], list(doc.get("notes", []))
    scanned = bool(doc.get("ocr"))
    if scanned:
        lines = _ocr_fix(lines)
    items = split_numbered(lines, ocr=scanned)
    recs = build_records(items) if items else []
    if len(recs) < 2:
        fb = build_records(split_by_verbs(lines))
        if len(fb) > len(recs):
            recs = fb
            notes.append("No question numbers were found, so questions were detected from action verbs (Define, Explain, ...). Please review the list.")
    if not recs:
        raise ValueError("I couldn't find any questions in this file. Make sure it is a question paper with numbered questions.")
    if len(recs) >= MAX_Q:
        notes.append(f"Only the first {MAX_Q} questions are shown.")
    if scanned:
        notes.insert(0, f"Read with OCR ({doc['ocr']['engine']}, {round(doc['ocr']['confidence'] * 100)}% confidence). "
                        "Scanned text can contain mistakes, so check each question before relying on its level.")
    notes.append("Levels are predicted from the wording (mainly the action verb) by a model trained on an AI-course question bank; "
                 "treat them as a strong first pass, not a final judgement.")
    return {"filename": doc["name"], "pages": doc["pages"], "questions": recs, "summary": summarize(recs), "notes": notes,
            "kind": doc["kind"], "ocr": doc.get("ocr"), "doc_id": doc["id"], "raw_text": doc["raw"][:6000]}


# ----------------------------------------------------------------- 4. Excel export
def to_xlsx(d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1D4ED8"), Font(bold=True, color="FFFFFF")

    def put(ws, row, values, header=False):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            if isinstance(v, str):
                c.data_type = "s"                       # never treat text as a formula (e.g. a question starting with "=")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if header:
                c.fill, c.font = head_fill, head_font

    ws = wb.active
    ws.title = "Questions"
    put(ws, 1, ["No.", "Section", "Question", "Marks", "Bloom level (predicted)", "Difficulty", "Confidence", "Paper's own tag", "Review"], True)
    for i, q in enumerate(d["questions"], 2):
        flags = []
        if q.get("low_confidence"):
            flags.append("Low confidence")
        if q.get("declared_level") and q["declared_level"] != q["level"]:
            flags.append("Differs from paper's tag")
        put(ws, i, [q["id"], q.get("section", ""), q["text"], q.get("marks"), q["level"], q["difficulty"],
                    q["level_confidence"], q.get("declared_level") or "", "; ".join(flags)])
        ws.cell(row=i, column=7).number_format = "0%"
    for col, w in zip("ABCDEFGHI", [8, 18, 80, 8, 22, 12, 12, 18, 26]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    s = d["summary"]
    ws2 = wb.create_sheet("Summary")
    put(ws2, 1, ["Bloom level", "Questions", "Share", "Marks"], True)
    for i, r in enumerate(s["by_level"], 2):
        put(ws2, i, [r["level"], r["count"], r["percent"] / 100, r["marks"] or None])
        ws2.cell(row=i, column=3).number_format = "0%"
    row = len(s["by_level"]) + 3
    put(ws2, row, ["Insights"], True)
    for t in s["insights"] + d.get("notes", []):
        row += 1
        put(ws2, row, [t])
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws2.row_dimensions[row].height = 32
    for col, w in zip("ABCD", [22, 12, 10, 10]):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
